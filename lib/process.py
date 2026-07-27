"""HR Excel processing engine — a faithful port of the Node backend's
routes/process.js (the single place all data manipulation happens in the
React + Express app).

Rules implemented, identical to the original:

1.  Header detection: column names are matched case/space/underscore/slash
    -insensitively. "DEPT" is treated as the departure-time column only when
    a separate "Department" column also exists.
2.  SPST normalization in place: WOP->WO, PHP->PH, X/X->X (WOP/WO -> WO...).
    Any row whose "Day" is Sunday has its SPST rewritten to WO, whatever it
    held before (DP, DP/WO, LWP, OD...).
3.  WO / PH / ABS rows — Sundays included, by rule 2: ARRV & DEPT blanked,
    WORK -> 0, OT Hours -> 0, and the row takes no distributed OT.
4.  Every other row: the row's "OT Hours" value is captured, then zeroed.
5.  ABS/DP & DP/ABS half-days: recorded punches kept; a worked span longer
    than 9.5 h is trimmed to 9.5 h minus 1..30 random minutes by moving
    DEPT (ABS/DP) or ARRV (DP/ABS); WORK recomputed.
6.  Rows whose status contains "DP" with both SHIFT IN & SHIFT OUT defined
    are collected per employee:
      - arrival is rewritten to shift-in + random 0..20 min,
      - departure snaps to shift-out unless the recorded punch-out is
        within +/-60 min of it (then the natural punch is kept),
      - a day qualifies for OT only when the status is exactly "DP" and the
        base worked span <= shift length + 10 min,
      - a base span longer than 9.5 h is trimmed to 9.5 h minus 0..30 min.
7.  Per employee, the monthly OT total (last positive "OT Hours" seen):
      - whole hours are split across qualifying days in 1..2 h chunks with
        jittered weights (distribute_ot),
      - a fractional remainder (e.g. the .5 of 12.5) rides on one random
        day as a 1 h + remainder chunk (or is added on its own if there is
        no whole hour to pair with),
      - each collected day is then written back: on OT days
        DEPT = ARRV + shift length + OT, otherwise the stored departure;
        WORK is recomputed for every collected day.
8.  Non-DP rows (PL, unknown, empty status with punches...): arrival is
    rewritten to shift-in + 0..20 min when a shift-in and any punch exist;
    WORK recomputed from the punches, trimming spans over 9.5 h (moving
    DEPT); WORK -> 0 when punches are incomplete.
9.  Times are written as Excel decimal day-fractions ("h:mm AM/PM" format
    for newly created cells), WORK uses the "[h]:mm" duration format —
    matching the SheetJS cell manipulation of the original.

Returns the processed workbook bytes plus the same stats the backend sends
in its X-Process-Stats header: employees, arrvFixed, deptFixed,
spstNormalized, workUpdated.
"""

from __future__ import annotations

import datetime as _dt
import io
import math
import random
import re

from openpyxl import Workbook, load_workbook

from .utils import (
    MIN_PER_DAY,
    js_parse_float,
    js_round,
    normalize_status,
    parse_time_to_minutes,
)

# Mirrors OT_CONFIG in routes/process.js
OT_CONFIG = {
    "minOtHours": 1,          # smallest OT chunk placed on a single day (hours)
    "maxOtHours": 2,          # largest OT chunk placed on a single day (hours)
    "arrvMaxLateMin": 20,     # arrival becomes shift-in + random 0..20 min
    "jitter": 0.4,            # weight jitter when splitting OT across days
    "otWorkedSlackMin": 10,   # OT-eligible if base worked <= shift len + 10 min
    "trimTargetMin": 9.5 * 60,   # cap for any non-OT worked span (9.5 h)
    "noOtTrimRandomMin": 30,     # random shave below the 9.5 h cap
    "naturalWindowMin": 60,   # keep a recorded punch within 60 min of shift
}


def _normalize_key(s) -> str:
    """normalizeKey(): lowercase, trim, strip spaces/underscores/slashes."""
    return re.sub(r"[\s_/]+", "", str(s).lower().strip())


def _rand_int(a: int, b: int) -> int:
    """randInt(): uniform integer in [a, b], both ends inclusive."""
    return random.randint(a, b)


def _wrap_day(m):
    """wrapDay(): fold any minute count into [0, 1440)."""
    return ((m % MIN_PER_DAY) + MIN_PER_DAY) % MIN_PER_DAY


# ── distributeOt(totalMin, n, minPerDay, maxPerDay, jitter) ───────────────
# Despite the parameter names, the call site passes whole HOURS: the total
# integer OT hours, split into 1..2-hour chunks over a random subset of the
# n qualifying days. Faithful port of the JS implementation.
def distribute_ot(total_min, n, min_per_day, max_per_day, jitter):
    alloc = [0] * n
    if n <= 0 or total_min <= 0:
        return alloc

    feasible = min(total_min, n * max_per_day)
    k_min = math.ceil(feasible / max_per_day)
    k_max = min(n, math.floor(feasible / min_per_day))

    if k_max < k_min:
        k = min(n, max(1, k_min))
        lo = math.floor(feasible / k)
    else:
        k = k_min + int(random.random() * (k_max - k_min + 1))
        if k < 1:
            k = 1
        lo = min_per_day

    # Fisher-Yates shuffle, then take the first k day indices
    order = list(range(n))
    for i in range(n - 1, 0, -1):
        j = int(random.random() * (i + 1))
        order[i], order[j] = order[j], order[i]
    chosen = order[:k]

    room = max_per_day - lo
    sel = [lo] * k
    rem = feasible - lo * k
    if rem > 0 and room > 0:
        weights = []
        wsum = 0.0
        for _ in range(k):
            w = 1 + (random.random() * 2 - 1) * jitter
            weights.append(w)
            wsum += w
        for i in range(k):
            add = js_round((rem * weights[i]) / wsum)
            if add > room:
                add = room
            if add < 0:
                add = 0
            sel[i] += add

    diff = feasible - sum(sel)
    guard = 0
    while diff != 0 and guard < 1_000_000:
        guard += 1
        i = int(random.random() * k)
        if diff > 0 and sel[i] < max_per_day:
            sel[i] += 1
            diff -= 1
        elif diff < 0 and sel[i] > lo:
            sel[i] -= 1
            diff += 1

    for i in range(k):
        alloc[chosen[i]] = sel[i]
    return alloc


# ── Workbook loading (xlsx/xlsm via openpyxl, legacy xls via xlrd, csv) ───
def _try_number(s: str):
    """SheetJS-style CSV type inference: a field that is a full number
    becomes a numeric cell (int when integral, so codes keep their text)."""
    if re.fullmatch(r"[+-]?(?:\d+\.?\d*|\.\d+)(?:[eE][+-]?\d+)?", s):
        f = float(s)
        if f.is_integer() and re.fullmatch(r"[+-]?\d+", s):
            return int(f)
        return f
    return None


def _csv_to_workbook(file_bytes: bytes) -> Workbook:
    import csv

    text = file_bytes.decode("utf-8-sig", errors="replace")
    wb = Workbook()
    ws = wb.active
    for r, row in enumerate(csv.reader(io.StringIO(text)), start=1):
        for c, raw in enumerate(row, start=1):
            val = raw.strip()
            if val == "":
                continue  # empty fields stay absent, like SheetJS
            num = _try_number(val)
            ws.cell(row=r, column=c, value=val if num is None else num)
    return wb


def _xls_to_workbook(file_bytes: bytes) -> Workbook:
    try:
        import xlrd  # legacy .xls only
    except ImportError as exc:  # pragma: no cover
        raise ValueError(
            "Legacy .xls support requires the 'xlrd' package "
            "(pip install xlrd), or re-save the file as .xlsx."
        ) from exc

    book = xlrd.open_workbook(file_contents=file_bytes)
    sh = book.sheet_by_index(0)
    wb = Workbook()
    ws = wb.active
    for r in range(sh.nrows):
        for c in range(sh.ncols):
            cell = sh.cell(r, c)
            if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                continue
            # date/time cells arrive as raw serial floats — exactly what
            # parse_time_to_minutes expects
            ws.cell(row=r + 1, column=c + 1, value=cell.value)
    return wb


def _load_workbook_any(file_bytes: bytes, filename: str) -> Workbook:
    name = (filename or "").lower()
    if name.endswith(".csv"):
        return _csv_to_workbook(file_bytes)
    if name.endswith(".xls"):
        return _xls_to_workbook(file_bytes)
    return load_workbook(io.BytesIO(file_bytes))


# ── The engine (router.post("/") body) ────────────────────────────────────
def process_workbook(file_bytes: bytes, filename: str = ""):
    """Process an attendance workbook. Returns (xlsx_bytes, stats_dict)."""
    wb = _load_workbook_any(file_bytes, filename)
    ws = wb[wb.sheetnames[0]]

    # if (!sheet || !sheet["!ref"]) — a truly empty sheet
    if ws.max_row is None or (
        ws.max_row == 1 and ws.max_column == 1 and ws.cell(1, 1).value is None
    ):
        raise ValueError("Excel file is empty or has no data rows.")

    header_row = ws.min_row
    max_row = ws.max_row
    min_col = ws.min_column
    max_col = ws.max_column

    # Header name -> column index
    col_of = {}
    for c in range(min_col, max_col + 1):
        hv = ws.cell(row=header_row, column=c).value
        if hv is not None and str(hv).strip() != "":
            col_of[_normalize_key(hv)] = c

    def pick(*names):
        for nm in names:
            if nm in col_of:
                return col_of[nm]
        return None

    # "Dept" is exit time only when a separate "Department" column exists
    has_department = "department" in col_of
    col = {
        "empKey": pick("employeecode", "code", "employeename", "name"),
        "spst": pick("spst", "status"),
        "shiftIn": pick("shiftin"),
        "shiftOut": pick("shiftout"),
        "arrv": pick("arrv", "arrival", "entry"),
        "dept": pick("dept") if has_department else None,
        "work": pick("work"),
        "ot": pick("othours"),
        "day": pick("day", "weekday"),
        "date": pick("date"),
    }

    if col["spst"] is None:
        raise ValueError("No SPST / Status column found in the sheet.")

    # ── Cell helpers (mirror the SheetJS ones) ────────────────────────────
    # openpyxl always returns a Cell object; "the cell does not exist" in
    # SheetJS terms means "it holds no value" here.
    def cell_at(R, C):
        return None if C is None else ws.cell(row=R, column=C)

    def read_min(R, C):
        c = cell_at(R, C)
        return parse_time_to_minutes(c.value) if c is not None else None

    def ensure_cell(R, C, z):
        # SheetJS creates a fresh {t:"n", z} cell; an existing cell keeps
        # its own number format.
        c = ws.cell(row=R, column=C)
        if c.value is None:
            c.number_format = z
        return c

    def write_time(cell, minutes):
        cell.value = _wrap_day(minutes) / MIN_PER_DAY
        # Cells that held plain strings default to "General" and would
        # display the raw fraction — upgrade those to a readable time
        # format (display-only; stored values match the original engine).
        if cell.number_format in (None, "", "General"):
            cell.number_format = "h:mm AM/PM"

    def write_work(R, minutes):
        if col["work"] is None:
            return
        worked = minutes
        if worked < 0:
            worked += MIN_PER_DAY
        worked = js_round(worked)
        cell = ensure_cell(R, col["work"], "General")
        if worked > 0:
            cell.value = worked / MIN_PER_DAY
            cell.number_format = "[h]:mm"
        else:
            cell.value = 0
            cell.number_format = "General"

    def zero_cell(cell):
        # JS: no-op when the cell is absent from the sheet
        if cell is None or cell.value is None:
            return
        cell.value = 0
        if not cell.number_format:
            cell.number_format = "General"

    def blank_cell(cell):
        if cell is None or cell.value is None:
            return
        cell.value = ""

    def is_sunday(R):
        """The "Day" column ("Sun", "Sunday") decides; a real date cell is
        the fallback when the sheet has no day-name column."""
        c = cell_at(R, col["day"])
        if c is not None and c.value is not None:
            return str(c.value).strip().lower().startswith("sun")
        c = cell_at(R, col["date"])
        if c is not None and isinstance(c.value, (_dt.datetime, _dt.date)):
            return c.value.weekday() == 6
        return False

    arrv_fixed = 0
    dept_fixed = 0
    spst_normalized = 0
    work_updated = 0

    employees: dict[str, dict] = {}  # insertion-ordered, like the JS Map

    # ── First pass: per-row normalization + collection ────────────────────
    for R in range(header_row + 1, max_row + 1):
        spst_cell = cell_at(R, col["spst"])
        # if (!spstCell) continue — a row whose SPST cell is absent is
        # skipped entirely
        if spst_cell is None or spst_cell.value is None:
            continue

        status = str(spst_cell.value or "").upper().strip()
        ns = normalize_status(status)
        sunday = is_sunday(R)

        # A Sunday is the weekly off whatever the raw status says (DP, DP/WO,
        # LWP, OD...), so its SPST is rewritten to WO.
        if sunday:
            ns = "WO"

        # 1) Normalize SPST in place
        if ns and ns != status:
            spst_cell.value = ns
            spst_normalized += 1

        ot_cell = cell_at(R, col["ot"])
        arrv_cell = cell_at(R, col["arrv"])
        dept_cell = cell_at(R, col["dept"])

        # 2) Sunday / off / holiday / absent: blank punches, zero WORK and OT.
        #    A Sunday is blanked whatever its status, so DP, DP/WO, LWP, OD...
        #    rows never show an arrival or a departure and never take OT.
        if ns in ("WO", "PH", "ABS"):
            blank_cell(arrv_cell)
            blank_cell(dept_cell)
            write_work(R, 0)
            zero_cell(ot_cell)
            continue

        # 3) Capture the row's OT Hours value, then zero the cell
        ot_val = js_parse_float(ot_cell.value) if ot_cell is not None else float("nan")
        zero_cell(ot_cell)

        si_min = read_min(R, col["shiftIn"])
        so_min = read_min(R, col["shiftOut"])

        # 4) Half-days: keep the punches, trim anything over 9.5 h
        if ns in ("ABS/DP", "DP/ABS"):
            a = read_min(R, col["arrv"])
            d = read_min(R, col["dept"])
            if a is not None and d is not None:
                w = d - a
                if w < 0:
                    w += MIN_PER_DAY
                if w > OT_CONFIG["trimTargetMin"]:
                    target = OT_CONFIG["trimTargetMin"] - _rand_int(
                        1, OT_CONFIG["noOtTrimRandomMin"]
                    )
                    if ns == "ABS/DP":
                        dc = cell_at(R, col["dept"])
                        if dc is not None:
                            write_time(dc, a + target)
                    else:
                        ac = cell_at(R, col["arrv"])
                        if ac is not None:
                            write_time(ac, d - target)
                    w = target
                write_work(R, w)
            else:
                write_work(R, 0)
            continue

        # 5) OT-managed rows: any status containing "DP" with a full shift
        eligible = ("DP" in ns) and si_min is not None and so_min is not None

        if eligible:
            if col["empKey"] is not None:
                kc = cell_at(R, col["empKey"])
                key = str((kc.value if kc is not None else None) or "").strip()
            else:
                key = "__all__"
            e = employees.get(key)
            if e is None:
                e = {"ot": 0.0, "days": []}
                employees[key] = e
            if ot_val > 0:  # NaN comparisons are False, like in JS
                e["ot"] = ot_val

            W = OT_CONFIG["naturalWindowMin"]
            orig_dep = read_min(R, col["dept"])
            shift_len = ((so_min - si_min) % MIN_PER_DAY + MIN_PER_DAY) % MIN_PER_DAY

            # arrival = shift-in + 0..20 random minutes
            arr = si_min + _rand_int(0, OT_CONFIG["arrvMaxLateMin"])
            if col["arrv"] is not None:
                write_time(ensure_cell(R, col["arrv"], "h:mm AM/PM"), arr)
            arrv_fixed += 1

            # departure: keep a natural punch near shift-out, else snap
            if orig_dep is not None and abs(orig_dep - so_min) <= W:
                dep = orig_dep
            else:
                dep = so_min

            base_worked = dep - arr
            if base_worked < 0:
                base_worked += MIN_PER_DAY

            # OT only on pure DP days without excessive recorded work
            ot_ok = ns == "DP" and base_worked <= shift_len + OT_CONFIG["otWorkedSlackMin"]

            if base_worked > OT_CONFIG["trimTargetMin"]:
                dep = (
                    arr
                    + OT_CONFIG["trimTargetMin"]
                    - _rand_int(0, OT_CONFIG["noOtTrimRandomMin"])
                )

            e["days"].append(
                {"R": R, "arr": arr, "dep": dep, "otOk": ot_ok, "shiftLen": shift_len}
            )
        else:
            # 6) Everything else (PL, unknown, empty status...)
            orig_arr = read_min(R, col["arrv"])
            orig_dep = read_min(R, col["dept"])

            a = orig_arr
            if si_min is not None and (orig_arr is not None or orig_dep is not None):
                a = si_min + _rand_int(0, OT_CONFIG["arrvMaxLateMin"])
                if col["arrv"] is not None:
                    write_time(ensure_cell(R, col["arrv"], "h:mm AM/PM"), a)
                arrv_fixed += 1

            d = orig_dep
            if a is not None and d is not None:
                w = d - a
                if w < 0:
                    w += MIN_PER_DAY

                if w > OT_CONFIG["trimTargetMin"]:
                    target = OT_CONFIG["trimTargetMin"] - _rand_int(
                        0, OT_CONFIG["noOtTrimRandomMin"]
                    )
                    dc = cell_at(R, col["dept"])
                    if dc is not None:
                        write_time(dc, a + target)
                    w = target
                write_work(R, w)
            else:
                write_work(R, 0)

    # ── Second pass: distribute OT and write departures / work ────────────
    for e in employees.values():
        ot_raw = e["ot"] or 0
        int_hours = math.floor(ot_raw)
        frac_min = js_round((ot_raw - int_hours) * 60)  # e.g. 0.5 h -> 30

        ot_days = [d for d in e["days"] if d["otOk"]]
        day_ot_min = [0] * len(ot_days)
        pool = list(range(len(ot_days)))

        # A fractional remainder rides on one random day as 1 h + frac
        half_placed = False
        if frac_min > 0 and int_hours >= 1 and len(pool) >= 1:
            half_idx = pool[random.randrange(len(pool))]
            day_ot_min[half_idx] = 60 + frac_min
            int_hours -= 1
            pool = [i for i in pool if i != half_idx]
            half_placed = True

        alloc = distribute_ot(
            int_hours,
            len(pool),
            OT_CONFIG["minOtHours"],
            OT_CONFIG["maxOtHours"],
            OT_CONFIG["jitter"],
        )
        for j in range(len(pool)):
            day_ot_min[pool[j]] += alloc[j] * 60

        if frac_min > 0 and not half_placed and len(ot_days) >= 1:
            day_ot_min[random.randrange(len(ot_days))] += frac_min

        for i, d in enumerate(ot_days):
            d["otMin"] = day_ot_min[i]

        for d in e["days"]:
            ot_min = d.get("otMin", 0) or 0

            arr = d["arr"]
            dep = arr + d["shiftLen"] + ot_min if ot_min > 0 else d["dep"]

            if col["dept"] is not None:
                write_time(ensure_cell(d["R"], col["dept"], "h:mm AM/PM"), dep)
                dept_fixed += 1

            worked = dep - arr
            if worked < 0:
                worked += MIN_PER_DAY
            write_work(d["R"], worked)
            work_updated += 1

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)

    stats = {
        "employees": len(employees),
        "arrvFixed": arrv_fixed,
        "deptFixed": dept_fixed,
        "spstNormalized": spst_normalized,
        "workUpdated": work_updated,
    }
    return out.getvalue(), stats
